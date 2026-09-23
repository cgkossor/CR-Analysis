"""Synthetic disintegration data with a known ground truth.

    python -m pipeline.disintegration.synthetic --input <dissolution.xlsx> \\
        --output outputs/synthetic/<name>_with_DT_SYNTHETIC.xlsx

Copies the workbook and adds two sheets: ``Disintegration`` (the data, in the
layout the loader expects) and ``Disintegration_Notes`` (a SYNTHETIC marker and
every true parameter). Because the truth is written down, the analysis can be
tested on whether it *recovers* it, not merely on whether it runs.

The data is tied to the workbook's own dissolution. Each design point is anchored
on its fitted Weibull time scale Td, which exists for every formulation, whereas
t80 is censored for most of the high-viscosity ones. Then, on the natural-log
scale::

    ln DT = ln Td + gamma[g] + delta[g] * zH - kappa * zA + eta + eps

    zH = (HPMC wt% - 40) / 20        zA = (API wt% - 35) / 25

The mechanism this encodes, which the analysis is expected to find:

* Within a grade, DT tracks the dissolution time scale one-for-one. That gives
  the high overall correlation.
* Across grades, the offset gamma *falls* as viscosity rises. The disintegration
  apparatus erodes a swollen gel mechanically, so a K100M matrix is gone well
  before its diffusion-controlled release is complete, while a K100LV matrix
  outlasts its release. **At identical composition, grade separates dissolution
  far more than it separates disintegration.**
* The HPMC lever on DT, delta, steepens with grade. A thicker gel layer resists
  mechanical erosion more.
* A higher drug load (a soluble API) leaves a more porous matrix that erodes
  sooner (kappa > 0).
* Replicate scatter grows with viscosity (heteroscedastic by grade). One
  discordant tablet is planted so the outlier test has something to find.
* Replicates past the test end are written as ``>1440`` (right-censored).
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from pipeline.analysis import run_analysis
from pipeline.disintegration import settings
from pipeline.io.load import load_database

#: Ground truth by grade. Grades not listed fall back to the nearest listed
#: viscosity, so the generator also runs on a workbook with other grades.
GAMMA: dict[str, float] = {"K100LV": 0.60, "K4M": -0.15, "K100M": -0.95}
DELTA: dict[str, float] = {"K100LV": 0.05, "K4M": 0.20, "K100M": 0.35}
REPLICATE_CV: dict[str, float] = {"K100LV": 0.06, "K4M": 0.08, "K100M": 0.12}
KAPPA = 0.15
ETA_SD = 0.05
HPMC_CENTRE, HPMC_SCALE = 40.0, 20.0
API_CENTRE, API_SCALE = 35.0, 25.0
OUTLIER_FACTOR = 2.0
TEST_END_MIN = 1440.0
_VISC = {"K100LV": 100.0, "K4M": 4_000.0, "K100M": 100_000.0}


def _nearest(grade: str, viscosity: float) -> str:
    if grade in GAMMA:
        return grade
    return min(_VISC, key=lambda g: abs(math.log10(_VISC[g]) - math.log10(max(viscosity, 1))))


def build_frames(source: Path, seed: int = settings.SYNTHETIC_SEED) -> tuple[
    pd.DataFrame, list[tuple[str, float, str]]
]:
    """(Disintegration sheet as a frame, truth rows) for ``source``."""
    analysis = run_analysis(load_database(source))
    points = analysis.design_points.copy()
    points["_order"] = points["viscosity_cp"]
    points = points.sort_values(["case", "_order"], kind="mergesort").reset_index(drop=True)
    ids = (
        analysis.replicates.groupby(["case", "grade"])["id"].first().to_dict()
    )

    rng = np.random.default_rng(seed)
    n_max = 4
    rows: list[dict[str, object]] = []
    planted: tuple[int, str, int] | None = None
    for p in points.itertuples():
        grade = str(p.grade)
        ref = _nearest(grade, float(p.viscosity_cp))
        ln_td = float(p.log10_td_mean) * math.log(10.0)
        z_h = (float(p.hpmc_wt) - HPMC_CENTRE) / HPMC_SCALE
        z_a = (float(p.api_wt) - API_CENTRE) / API_SCALE
        mu = ln_td + GAMMA[ref] + DELTA[ref] * z_h - KAPPA * z_a + rng.normal(0.0, ETA_SD)
        n = int(rng.choice([3, 4]))
        sigma = math.sqrt(math.log1p(REPLICATE_CV[ref] ** 2))
        reps = np.exp(mu + rng.normal(0.0, sigma, n)) * 60.0  # minutes

        # One discordant tablet, in the first low-viscosity formulation run in
        # quadruplicate: the tightest replicates, so it is the fairest test.
        if planted is None and ref == "K100LV" and n == 4:
            reps[2] *= OUTLIER_FACTOR
            planted = (int(p.case), grade, 3)

        row: dict[str, object] = {
            "ID": ids.get((p.case, p.grade), f"{p.api}_{grade}_C{int(p.case):02d}"),
            "Case": int(p.case),
            "API": str(p.api),
            "HPMC Grade": grade,
            "API [wt%]": float(p.api_wt),
            "HPMC [wt%]": float(p.hpmc_wt),
            "Lactose [wt%]": float(p.lactose_wt),
        }
        for k in range(n_max):
            if k >= n:
                row[f"DT_{k + 1} [min]"] = None
                continue
            minutes = round(float(reps[k]), 1)
            row[f"DT_{k + 1} [min]"] = (
                f">{TEST_END_MIN:g}" if minutes >= TEST_END_MIN else minutes
            )
        row["Test_end [min]"] = TEST_END_MIN
        rows.append(row)

    truth: list[tuple[str, float, str]] = []
    for g in GAMMA:
        truth.append((f"gamma[{g}]", GAMMA[g], "ln-scale offset of DT from Td"))
    for g in DELTA:
        truth.append((f"delta[{g}]", DELTA[g], "HPMC slope on ln DT, per 20 wt%"))
    for g in REPLICATE_CV:
        truth.append((f"cv[{g}]", REPLICATE_CV[g], "replicate CV"))
    truth += [
        ("kappa", KAPPA, "API slope (subtracted) on ln DT, per 25 wt%"),
        ("eta_sd", ETA_SD, "formulation lack-of-fit SD, ln scale"),
        ("hpmc_centre", HPMC_CENTRE, "wt%"),
        ("hpmc_scale", HPMC_SCALE, "wt%"),
        ("api_centre", API_CENTRE, "wt%"),
        ("api_scale", API_SCALE, "wt%"),
        ("outlier_factor", OUTLIER_FACTOR, "multiplier on the planted replicate"),
        ("test_end_min", TEST_END_MIN, "min"),
        ("seed", float(seed), ""),
    ]
    if planted is not None:
        truth += [
            ("outlier_case", float(planted[0]), "case of the planted outlier"),
            ("outlier_replicate", float(planted[2]), f"replicate, grade {planted[1]}"),
        ]
    return pd.DataFrame(rows), truth


def generate(source: str | Path, output: str | Path, seed: int = settings.SYNTHETIC_SEED) -> Path:
    """Write ``output``: ``source`` plus the two disintegration sheets."""
    src, dst = Path(source), Path(output)
    frame, truth = build_frames(src, seed)
    wb = openpyxl.load_workbook(src)
    for name in ("Disintegration", "Disintegration_Notes"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("Disintegration")
    ws.append(list(frame.columns))
    for rec in frame.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and math.isnan(v)) else v for v in rec])

    notes = wb.create_sheet("Disintegration_Notes")
    notes.append(["SYNTHETIC PLACEHOLDER DATA — NOT EXPERIMENTAL"])
    notes.append(["Generated by pipeline.disintegration.synthetic; model in its docstring."])
    notes.append(["ln DT = ln Td + gamma[g] + delta[g]*zH - kappa*zA + eta + eps"])
    notes.append([])
    notes.append(["parameter", "value", "meaning"])
    for key, value, meaning in truth:
        notes.append([key, value, meaning])

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m pipeline.disintegration.synthetic")
    parser.add_argument("--input", required=True, help="dissolution workbook to extend")
    parser.add_argument(
        "--output",
        default=str(Path("outputs") / "synthetic" / "CR_with_DT_SYNTHETIC.xlsx"),
        help="where to write the extended copy (the input is never modified)",
    )
    parser.add_argument("--seed", type=int, default=settings.SYNTHETIC_SEED)
    args = parser.parse_args(argv)
    out = generate(args.input, args.output, args.seed)
    print(f"Synthetic disintegration workbook -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
