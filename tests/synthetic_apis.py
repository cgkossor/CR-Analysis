"""Synthetic second (and third) APIs for the multi-API tests, never for analysis.

Each is the placeholder workbook with its API renamed and every profile slowed
or sped up in time, roughly how a less or more soluble drug shifts release
from the same matrix. The source workbook already declares itself synthetic
in its Notes sheet, and the copies keep that sheet, so every output drawn from
them carries the provenance banner (G2).
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl


def make_api_workbook(source: Path, dest: Path, api: str, time_scale: float) -> Path:
    """Copy ``source`` as API ``api`` with release stretched in time by ``time_scale``.

    ``time_scale`` > 1 slows release (the concentration reached at time t is the
    original one at t / time_scale); < 1 speeds it up.
    """
    wb = openpyxl.load_workbook(source)
    ws = wb["Dissolution"]
    header = [c.value for c in ws[1]]
    col = {str(h): i for i, h in enumerate(header) if h is not None}
    id_i, api_i, time_i = col["ID"], col["API"], col["Min_1"]
    conc = [i for h, i in col.items() if h.lower().startswith("conc_")]

    rows: dict[str, list[int]] = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, id_i + 1).value
        if rid is None:
            continue
        rows.setdefault(str(rid), []).append(r)

    for rid, rs in rows.items():
        t = np.array([float(ws.cell(r, time_i + 1).value) for r in rs])
        for ci in conc:
            y = np.array([float(ws.cell(r, ci + 1).value or 0.0) for r in rs])
            warped = np.interp(t / time_scale, t, y)
            for r, v in zip(rs, warped, strict=True):
                ws.cell(r, ci + 1).value = round(float(v), 4)
        for r in rs:
            ws.cell(r, api_i + 1).value = api
            ws.cell(r, id_i + 1).value = rid.replace("API_1", api)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
