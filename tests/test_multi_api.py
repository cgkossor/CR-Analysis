"""Several APIs: one workbook each, never merged, selectable on the dashboard."""

from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest
from synthetic_apis import make_api_workbook

from pipeline.io.api_props import load_api_props
from pipeline.io.load import load_database
from pipeline.io.schema import SchemaError
from pipeline.run import main

ROOT = Path(__file__).resolve().parents[1]
DASHBOARD = ROOT / "dashboard"
NODE = shutil.which("node")


def _workbook() -> Path:
    matches = sorted(ROOT.glob("*.xlsx"))
    if not matches:
        pytest.skip("no database workbook present")
    return matches[0]


def test_a_workbook_with_two_apis_is_refused(tmp_path: Path) -> None:
    """Merging two APIs on (case, grade) would average them silently."""
    wb = openpyxl.load_workbook(_workbook())
    ws = wb["Dissolution"]
    header = [c.value for c in ws[1]]
    api_col = header.index("API") + 1
    for r in range(2, ws.max_row // 2):
        ws.cell(r, api_col).value = "API_X"
    mixed = tmp_path / "mixed.xlsx"
    wb.save(mixed)
    with pytest.raises(SchemaError) as err:
        load_database(mixed)
    assert err.value.code == 33


def test_api_properties_are_validated(tmp_path: Path) -> None:
    good = tmp_path / "apis.csv"
    good.write_text("api,solubility_mg_ml,solubility_class\nA,33,high\nB,0.05,\n",
                    encoding="utf-8")
    props = load_api_props(good)
    assert props["A"] == {"solubility_mg_ml": 33.0, "solubility_class": "high"}
    assert props["B"]["solubility_class"] is None

    for body in ("api,solubility_mg_ml\nA,abc\n", "api,solubility_mg_ml\nA,-1\n",
                 "api,solubility_mg_ml,solubility_class\nA,1,medium\n",
                 "api,solubility_mg_ml\nA,1\nA,2\n", "api\nA\n"):
        bad = tmp_path / "bad.csv"
        bad.write_text(body, encoding="utf-8")
        with pytest.raises(SchemaError) as err:
            load_api_props(bad)
        assert err.value.code == 34


@pytest.fixture(scope="module")
def two_api_run(tmp_path_factory: pytest.TempPathFactory) -> Path:
    tmp = tmp_path_factory.mktemp("multi")
    second = make_api_workbook(_workbook(), tmp / "api_b.xlsx", "API_B", 1.6)
    apis = tmp / "apis.csv"
    apis.write_text("api,solubility_mg_ml,solubility_class\nAPI_1,33,high\nAPI_B,0.05,low\n",
                    encoding="utf-8")
    dash = tmp / "dashboard"
    dash.mkdir()
    for f in DASHBOARD.iterdir():
        if f.suffix in (".html", ".js", ".css") and f.name != "data.js":
            shutil.copy(f, dash / f.name)
    code = main([
        "--input", str(_workbook()), "--input", str(second), "--apis", str(apis),
        "--outputs", str(tmp / "outputs"), "--dashboard", str(dash), "--skip-figures",
    ])
    assert code == 0
    return tmp


def test_each_api_gets_its_own_outputs(two_api_run: Path) -> None:
    for api in ("API_1", "API_B"):
        reports = two_api_run / "outputs" / api / "reports"
        assert (reports / "diagnostics.md").exists(), api
        assert (reports / "storyline.md").exists(), api


def test_data_js_carries_every_api(two_api_run: Path) -> None:
    text = (two_api_run / "dashboard" / "data.js").read_text(encoding="utf-8")
    by_api = json.loads(text.split("window.CR_DATA_BY_API = ", 1)[1].split(";\n", 1)[0])
    api_set = json.loads(text.split("window.CR_API_SET = ", 1)[1].split(";\n", 1)[0])
    assert set(by_api) == {"API_1", "API_B"}
    assert api_set["order"] == ["API_1", "API_B"]
    assert api_set["props"]["API_B"]["solubility_class"] == "low"
    for api, payload in by_api.items():
        assert payload["quality"]["apis"] == [api], "a payload mixes APIs"


def test_disintegration_files_must_pair_with_inputs(tmp_path: Path) -> None:
    code = main(["--input", "a.xlsx", "--input", "b.xlsx", "--disintegration", "dt.xlsx",
                 "--outputs", str(tmp_path)])
    assert code == 2


def test_the_same_api_twice_is_refused(tmp_path: Path) -> None:
    code = main(["--input", str(_workbook()), "--input", str(_workbook()),
                 "--outputs", str(tmp_path / "o"), "--dashboard", str(tmp_path / "d"),
                 "--skip-figures"])
    assert code == 2


_SELECT = r"""
const { JSDOM, VirtualConsole } = require("jsdom");
const fs = require("fs"), path = require("path");
const dash = process.argv[1], hash = process.argv[2];
const errors = [];
const vc = new VirtualConsole();
vc.on("jsdomError", e => errors.push("jsdomError: " + e.message));
vc.on("error", (...a) => errors.push("console.error: " + a.join(" ")));
const dom = new JSDOM(fs.readFileSync(path.join(dash, "index.html"), "utf8"),
  { runScripts: "dangerously", virtualConsole: vc, url: "http://localhost/index.html" + hash });
const w = dom.window;
for (const f of ["data.js", "model.js", "doe.js", "formulator.js", "disintegration.js",
                 "figures.js", "manuscript.js", "audit.js", "app.js"]) {
  w.eval(fs.readFileSync(path.join(dash, f), "utf8"));
}
w.document.dispatchEvent(new w.Event("DOMContentLoaded"));
const d = w.document, sel = d.getElementById("api-select");
process.stdout.write(JSON.stringify({
  options: sel ? Array.from(sel.options).map(o => o.value) : [],
  selected: sel ? sel.value : null,
  subtitle: d.getElementById("subtitle").textContent,
  hash: w.location.hash,
  errors,
}));
"""


def _render(dash: Path, hash_: str) -> dict:
    if NODE is None:
        pytest.skip("node is not installed")
    probe = subprocess.run([NODE, "-e", "require.resolve('jsdom')"], capture_output=True,
                           check=False, cwd=ROOT)
    if probe.returncode != 0:
        pytest.skip("jsdom not installed")
    result = subprocess.run([NODE, "-e", _SELECT, str(dash), hash_], capture_output=True,
                            text=True, timeout=300, check=False, cwd=ROOT)
    assert result.returncode == 0, result.stderr[:2000]
    return json.loads(result.stdout)


def test_dashboard_offers_and_honours_the_api_choice(two_api_run: Path) -> None:
    dash = two_api_run / "dashboard"
    first = _render(dash, "")
    assert first["errors"] == []
    assert first["options"] == ["API_1", "API_B"]
    assert first["selected"] == "API_1" and first["subtitle"].startswith("API_1")

    second = _render(dash, "#manuscripts/q1@API_B")
    assert second["errors"] == []
    assert second["selected"] == "API_B" and second["subtitle"].startswith("API_B")
    assert second["hash"] == "#manuscripts/q1@API_B", "navigation dropped the API"


def test_api_names_sharing_a_folder_are_refused(tmp_path: Path) -> None:
    """'API 1' and 'API_1' would both write outputs/API_1; the second must not run."""
    other = make_api_workbook(_workbook(), tmp_path / "spaced.xlsx", "API 1", 1.0)
    out = tmp_path / "outputs"
    code = main(["--input", str(_workbook()), "--input", str(other),
                 "--outputs", str(out), "--dashboard", str(tmp_path / "d"), "--skip-figures"])
    assert code == 2
    assert not (tmp_path / "d" / "data.js").exists(), "data.js written after a failed input"


def test_paper_figures_span_every_api(tmp_path: Path) -> None:
    """The curated main figures are drawn once, across all APIs in the run."""
    from pipeline.analysis import run_analysis
    from pipeline.manuscript import build
    from pipeline.paper import PaperInput, headline_response, render_paper
    from pipeline.run import _stress

    second = make_api_workbook(_workbook(), tmp_path / "api_b.xlsx", "API_B", 1.6)
    inputs = []
    for api, src in (("API_1", _workbook()), ("API_B", second)):
        a = run_analysis(load_database(src))
        inputs.append(PaperInput(api, a, _stress(a), None, build(a, None)))
    records = render_paper(inputs, tmp_path / "paper")
    ids = [r.id for r in records]
    # No disintegration data here, so Fig 7 is left out rather than drawn empty.
    assert ids == ["Scheme1", "Fig1", "Fig2", "Fig3", "Fig4", "Fig5", "Fig6", "Fig6A",
                   "Fig8", "Fig8A"]
    # t50 is too sparse for API_B, so Fig 4 is already on MDT and needs no Fig 4A.
    assert headline_response(inputs) == "mdt_h"
    assert headline_response(inputs) is not None
    for name in ("captions.json", "captions.md", "Table1_design.csv", "supplementary.md"):
        assert (tmp_path / "paper" / name).exists(), name
    caps = json.loads((tmp_path / "paper" / "captions.json").read_text(encoding="utf-8"))
    assert any("API_1" in c["caption"] and "API_B" in c["caption"] for c in caps)
    # Every main figure says why it matters, beside its caption.
    assert all(c["context"] for c in caps)
    assert "Why it matters." in (tmp_path / "paper" / "captions.md").read_text(encoding="utf-8")


def _blank_replicate(source: Path, dest: Path, column: str, ids: int) -> Path:
    """Copy ``source`` with ``column`` emptied for the first ``ids`` IDs."""
    wb = openpyxl.load_workbook(source)
    ws = wb["Dissolution"]
    header = [c.value for c in ws[1]]
    col = header.index(column) + 1
    id_col = header.index("ID") + 1
    seen: list[object] = []
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, id_col).value
        if rid not in seen:
            seen.append(rid)
        if len(seen) <= ids:
            ws.cell(r, col).value = None
    wb.save(dest)
    return dest


def test_a_blank_replicate_is_dropped_not_fatal(tmp_path: Path) -> None:
    """A vessel with no readings used to crash the metrics table (KeyError 'weibull')."""
    header = [c.value for c in openpyxl.load_workbook(_workbook())["Dissolution"][1]]
    conc = next(h for h in header if str(h).lower().startswith("conc_"))
    bad = _blank_replicate(_workbook(), tmp_path / "blank_rep.xlsx", conc, 2)
    with pytest.warns(UserWarning, match="no usable reading"):
        db = load_database(bad)
    assert db.profiles.groupby(["id", "replicate"]).size().min() > 0
    code = main(["--input", str(bad), "--outputs", str(tmp_path / "o"),
                 "--dashboard", str(tmp_path / "d"), "--skip-figures"])
    assert code == 0


def test_mdt_alternative_accompanies_a_t50_surface(tmp_path: Path) -> None:
    """When Fig 4 is on t50, Fig 4A redraws it on MDT, which every blend has."""
    from pipeline.analysis import run_analysis
    from pipeline.manuscript import build
    from pipeline.paper import PaperInput, headline_response, render_paper
    from pipeline.run import _stress

    a = run_analysis(load_database(_workbook()))
    inputs = [PaperInput("API_1", a, _stress(a), None, build(a, None))]
    assert headline_response(inputs) == "t50"
    records = {r.id: r for r in render_paper(inputs, tmp_path / "paper")}
    alt = records["Fig4A"]
    assert alt.rank == records["Fig4"].rank
    assert "mean dissolution time" in alt.caption and "lower bound" in alt.caption
    assert (tmp_path / "paper" / alt.file).exists()


def test_run_plan_lists_exactly_the_recommended_runs(tmp_path: Path) -> None:
    """Fig 8A names every run of the recommended reduced design, and no other."""
    from pipeline.analysis import run_analysis
    from pipeline.manuscript import build
    from pipeline.paper import PaperInput, run_plan
    from pipeline.run import _stress

    a = run_analysis(load_database(_workbook()))
    p = PaperInput("API_1", a, _stress(a), None, build(a, None))
    plan = run_plan(p)
    assert plan is not None and p.stress.recommended is not None
    listed = {(c, g) for g, cs in plan.items() for c in cs}
    assert listed == {(int(c), str(g)) for c, g in p.stress.recommended.selected}
